# imports
import streamlit as st
from io import StringIO, BytesIO
from Bio import SeqIO
from seleniumbase import SB
import pandas as pd
from openpyxl import load_workbook
import time

# Block for Streamlit cloud
import os
from seleniumbase.core import browser_launcher
os.makedirs("/tmp/seleniumbase_drivers", exist_ok=True)
browser_launcher.override_driver_dir("/tmp/seleniumbase_drivers")

st.set_page_config(page_title="Gibson Automator", page_icon="🧬", layout="centered")

st.markdown(
    """
    <style>
    .stApp {
        background: linear-gradient(135deg, #0E1117 0%, #1A2E35 100%);
    }
    
    [data-testid="stAppViewContainer"] {
        background-color: transparent;
    }
    </style>
    """,
    unsafe_allow_html=True
)

# Setting up session memory
if "excel_file" not in st.session_state:
    st.session_state.excel_file = None
if "df" not in st.session_state:
    st.session_state.df = None

st.title("🧬 Gibson Assembly Automator")

with st.expander("📌 Benchling Export Instructions", expanded=True):
    st.info("""
    **1.** Paste your DNA sequences and annotate each fragment.\n
    **2.** Click the **ⓘ** icon located at the very bottom of the right-sided panel.\n
    **3.** Scroll down, export the data as **.gb**, and click **Download**.
    """)
    st.caption("Having a hard time getting it right? Ask Guy for help.")

# File Uploader
uploaded_file = st.file_uploader("Upload .gb file", type=["gb", "gbk"])

if uploaded_file is not None:
    if st.button("Generate Primers", type="primary", use_container_width=True):
        start_time = time.time()

        
        # --- FILE READING BLOCK ---
        stringio = StringIO(uploaded_file.getvalue().decode("utf-8"))
        record = SeqIO.read(stringio, "genbank")
        
        plasmid_seq = record.seq
        feature_dict = {}
        for feature in record.features:
            label = feature.qualifiers.get('label', [feature.type])[0]
            feature_sequence = feature.location.extract(plasmid_seq)
            feature_dict[label] = str(feature_sequence)

        if len(feature_dict) == 0:
            st.error("No annotated fragments found in the uploaded file. Please check Benchling and try again.")
            st.stop()

        with st.status("Automating Primer Design...\nDo not close the browser!", expanded=True) as status:
            
            # --- BROWSER AUTOMATION BLOCK ---
            with SB(uc=True, headless=True) as sb:
                sb.driver.get("https://nebuilder.neb.com/#!/")
                st.write("\nGoing to", sb.driver.title,"...\n")

                for label, seq in feature_dict.items():
                    sb.wait_for_element_clickable("//div[@class='btn btn-link-cta plusfrag']").click()
                    sb.click("//label[contains(normalize-space(), 'Paste Sequence')]")
                    sb.type("#target", seq)
                    sb.click("#parseseq")
                    sb.type("#fragname", label)
                    sb.click("div[ng-click='addNewFragment()']")

                # Extract primer table from NEBuilder
                primer_table = sb.driver.find_element("xpath", "//table[@class='table table-condensed table-striped  oligoTable']")
                rows = primer_table.find_elements("tag name", "tr")

                oligo_dict = {}
                for row in rows[1:]:
                    cols = row.find_elements("tag name", "td")
                    if len(cols) >= 2:
                        raw_name = cols[0].text
                        sequence = cols[1].text
                        fragment_name = raw_name.replace("_fwd", "").replace("_rev", "").strip()
                        
                        if fragment_name not in oligo_dict:
                            oligo_dict[fragment_name] = (sequence,)
                        else:
                            existing_fwd_seq = oligo_dict[fragment_name][0]
                            oligo_dict[fragment_name] = (existing_fwd_seq, sequence)

                fragments = list(oligo_dict.keys())
                seqs = list(oligo_dict.values())
                col_names = ['annealing_temp', 'tm_fwd', 'tm_rev', 'tm_diff', 'frag_name', 'fwd_seq', 'rev_seq']
                data = []

                sb.driver.get("https://tmcalculator.neb.com/#!/main")
                st.write("Going to", sb.driver.title,"...\n")

                for i, seq in enumerate(seqs):
                    status.update(label=f"Processing fragment {fragments[i]} ({i+1}/{len(feature_dict)})...")
                    data_row = []
                    sb.type("#p1", seq[0])
                    sb.type("#p2", seq[1])
                    
                    sb.sleep(0.5)

                    ta_text = sb.get_text("h2") 
                    ta = ta_text.split(" ")[0]
                    data_row.append(ta)
                    
                    strong_tags = sb.find_elements("strong")
                    tm1_text = strong_tags[4].text
                    tm2_text = strong_tags[7].text
                    tm1 = int(tm1_text.split(" ")[1].split("°")[0])
                    tm2 = int(tm2_text.split(" ")[1].split("°")[0])

                    data_row.append(tm1)
                    data_row.append(tm2)

                    diff = abs(tm1 - tm2) <= 5
                    data_row.append(diff)

                    data_row.append(fragments[i])
                    data_row.append(seq[0])
                    data_row.append(seq[1])

                    data.append(data_row)
            # --- END OF BROWSER SESSION ---

            # --- DATA PARSING & EXCEL GENERATION bLOCK ---
            df = pd.DataFrame(data)
            df.columns = col_names
            df.set_index('frag_name', inplace=True)

            st.session_state.df = df

            wb = load_workbook('files/template.xlsx')
            ws = wb.active
            current_row = 6 

            for frag_name, row_data in df.iterrows():
                fwd_seq = row_data['fwd_seq']
                rev_seq = row_data['rev_seq']
                
                ws.cell(row=current_row, column=1, value=f"{frag_name}_fwd")
                ws.cell(row=current_row, column=2, value=fwd_seq)
                ws.cell(row=current_row, column=5, value="0.025")
                ws.cell(row=current_row, column=6, value="Desalt")
                ws.cell(row=current_row, column=7, value="In Solution (water)")
                ws.cell(row=current_row, column=8, value="100")
                current_row += 1 
                
                ws.cell(row=current_row, column=1, value=f"{frag_name}_rev")
                ws.cell(row=current_row, column=2, value=rev_seq)
                ws.cell(row=current_row, column=5, value="0.025")
                ws.cell(row=current_row, column=6, value="Desalt")
                ws.cell(row=current_row, column=7, value="In Solution (water)")
                ws.cell(row=current_row, column=8, value="100")
                current_row += 1

            virtual_workbook = BytesIO()
            wb.save(virtual_workbook)
            virtual_workbook.seek(0)
            st.session_state.excel_file = virtual_workbook.getvalue()
            
            elapsed_time = time.time() - start_time
            minutes, seconds = divmod(int(elapsed_time), 60)
            time_str = f"{minutes}m {seconds}s" if minutes > 0 else f"{elapsed_time:.1f}s"
            status.update(label=f"Automation Complete in {time_str}! click to expand ⬇", state="complete", expanded=False)
    
    # --- UI DISPLAY BLOCK - primers summary, warnings ---
    if st.session_state.df is not None:
        st.divider()
        st.write("### 🔬 Primer Overview")
        
        st.dataframe(
            st.session_state.df.style.map(
                lambda x: 'background-color: rgba(255, 75, 75, 0.3)' if x is False else '', 
                subset=['tm_diff']
            ),
            use_container_width=True
        )
        
        warnings_found = False
        for frag_name, row_data in st.session_state.df.iterrows():
            if row_data['tm_diff'] == False:
                st.warning(f'Tm difference of **{frag_name}** is greater than 5. Please check on that.', icon="⚠️")
                warnings_found = True
        
        if not warnings_found:
            st.success("All Tm differences are within the optimal range (≤ 5°C).", icon="✅")

    # Download button
    if st.session_state.excel_file is not None:
        st.divider()
        st.success("✅ Sequence parsed and primers generated successfully.")
        
        st.download_button(
            label="📥 Download Sigma Order Form",
            data=st.session_state.excel_file,
            file_name=f"{uploaded_file.name.split('.')[0]}_sigma_order.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary"
        )